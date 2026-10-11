#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
=============================================================================
ClickAI Product Mapper - Script chuẩn hóa & ánh xạ dữ liệu sản phẩm
=============================================================================
Tự động đọc file (.xlsx, .xls, .csv), ánh xạ các trường thông tin về đúng
16 cột chuẩn ClickAI, điền giá trị mặc định và xuất file kết quả.
=============================================================================
"""

import os
import sys
import re
import json
import shutil
import argparse
import unicodedata
import xml.etree.ElementTree as ET
from typing import Union, Optional, Tuple, List, Dict
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Tắt cảnh báo downcasting của pandas trong các phiên bản mới
try:
    pd.set_option('future.no_silent_downcasting', True)
except Exception:
    pass


# 16 Cột tiêu chuẩn của ClickAI theo đúng thứ tự
CLICKAI_COLUMNS = [
    "Tên sản phẩm",
    "Tùy chọn giá",
    "Mã SKU",
    "Danh mục",
    "Giá bán",
    "Giá vốn",
    "Tồn kho",
    "Đơn vị tính",
    "Tiền tệ",
    "Thuế (%)",
    "Chi phí trực tiếp",
    "Chu kỳ thanh toán",
    "Xuất xứ",
    "Mô tả",
    "Link hình ảnh",
    "Link sản phẩm"
]

# Giá trị mặc định nếu trường dữ liệu bị trống
DEFAULT_VALUES = {
    "Tiền tệ": "VND",
    "Chu kỳ thanh toán": "Một lần",
    "Thuế (%)": 0,
    "Chi phí trực tiếp": 0,
    "Tồn kho": 0
}

# Từ điển ánh xạ từ khóa (Synonyms & Aliases)
# Hỗ trợ cả tiếng Việt có dấu, không dấu và tiếng Anh phổ biến từ các sàn TMĐT (Shopee, Lazada, TikTok, KiotViet, Sapo, Haravan, WooCommerce, Shopify)
SYNONYMS_MAP = {
    "Tên sản phẩm": [
        "ten san pham", "ten sp", "ten hang hoa", "ten hang", "ten", "tieu de", 
        "product name", "product_name", "product title", "product_title", "title", 
        "item name", "item_name", "name", "product", "g:title", "post_title", "headline",
        "ten vat tu", "vat tu", "ten hang hoa vat tu"
    ],
    "Tùy chọn giá": [
        "tuy chon gia", "bien the", "loai gia", "phan loai", "option", 
        "variant", "price option", "price_option", "variation", "variant name", 
        "thuoc tinh", "kich thuoc", "mau sac", "size", "color"
    ],
    "Mã SKU": [
        "ma sku", "sku", "ma san pham", "ma sp", "ma hang", "ma hang hoa", 
        "product code", "product_code", "item code", "code", "barcode", 
        "ma vach", "parent sku", "variant sku", "g:id", "id", "item_group_id",
        "ma vat tu", "ma vt"
    ],
    "Danh mục": [
        "danh muc", "nhom san pham", "nhom hang", "chuyen muc", "loai hang", 
        "category", "categories", "product category", "collection", "product type", 
        "nganh hang", "g:product_type", "g:google_product_category"
    ],
    "Giá bán": [
        "gia ban", "gia", "don gia", "gia niem yet", "gia ban le", "gia hien tai", 
        "price", "sale price", "sale_price", "regular price", "regular_price", 
        "selling price", "retail price", "unit price", "g:price", "g:sale_price"
    ],
    "Giá vốn": [
        "gia von", "gia nhap", "chi phi von", "cost", "cost price",
        "cost_price", "purchase price"
    ],
    "Tồn kho": [
        "ton kho", "so luong ton", "so luong", "sl ton", "sl", "kho", "kho hang", 
        "stock", "quantity", "qty", "inventory", "available quantity", "stock quantity", "g:quantity",
        "tong cong", "ton kha dung", "kha dung", "so luong kha dung"
    ],
    "Đơn vị tính": [
        "don vi tinh", "dvt", "don vi", "uom", "unit", "unit of measure", "pack"
    ],
    "Tiền tệ": [
        "tien te", "loai tien", "currency", "currency code", "curr"
    ],
    "Thuế (%)": [
        "thue (%)", "thue", "vat", "vat (%)", "tax", "tax rate", "tax (%)", 
        "thue suat", "thue vat"
    ],
    "Chi phí trực tiếp": [
        "chi phi truc tiep", "chi phi", "direct cost", "direct costs", "extra cost"
    ],
    "Chu kỳ thanh toán": [
        "chu ky thanh toan", "ky thanh toan", "billing cycle", "payment cycle", 
        "recurring period", "chu ky"
    ],
    "Xuất xứ": [
        "xuat xu", "noi san xuat", "quoc gia", "origin", "country of origin", 
        "made in", "brand origin", "g:country"
    ],
    "Mô tả": [
        "mo ta", "mo ta san pham", "chi tiet", "noi dung", "mo ta ngan", 
        "description", "short description", "product description", "details", 
        "body html", "content", "g:description", "post_content", "post_excerpt"
    ],
    "Link hình ảnh": [
        "link hinh anh", "hinh anh", "anh", "anh san pham", "link anh", 
        "image", "image url", "image_url", "images", "photo", "photos", 
        "image link", "picture", "thumbnail", "g:image_link", "g:additional_image_link"
    ],
    "Link sản phẩm": [
        "link san pham", "duong dan", "lien ket", "url", "product url", 
        "product_url", "link", "handle", "product link", "web link", "g:link", "permalink"
    ]
}


def remove_accents(input_str: str) -> str:
    """Loại bỏ dấu tiếng Việt, ký tự đặc biệt, chuẩn hóa về chữ thường."""
    if not isinstance(input_str, str):
        input_str = str(input_str)
    # Loại bỏ non-breaking space và ký tự ẩn
    input_str = input_str.replace('\xa0', ' ').replace('\u200b', '').strip()
    # Chuẩn hóa unicode và bỏ dấu
    nfkd_form = unicodedata.normalize('NFKD', input_str)
    without_accents = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    # Thay đ/Đ thành d/D
    without_accents = without_accents.replace('đ', 'd').replace('Đ', 'D')
    # Chuyển về chữ thường, thay ký tự đặc biệt thành dấu cách
    cleaned = re.sub(r'[^a-zA-Z0-9]', ' ', without_accents.lower())
    # Bỏ khoảng trắng thừa
    return " ".join(cleaned.split())


def count_matched_product_columns(column_names: list) -> int:
    """Đếm số lượng cột khớp với từ điển sản phẩm ClickAI để nhận diện header/sheet."""
    normalized_names = []
    for c in column_names:
        if pd.notna(c) and str(c).strip():
            s = str(c).strip()
            # Bỏ qua các giá trị là URL, số thuần túy hoặc quá dài (chắc chắn là hàng dữ liệu)
            if s.startswith("http://") or s.startswith("https://") or s.isdigit() or len(s) > 60:
                continue
            normalized_names.append(remove_accents(s))

    if not normalized_names:
        return 0
    matched_targets = set()
    for col_norm in normalized_names:
        for target_col, aliases in SYNONYMS_MAP.items():
            if target_col in matched_targets:
                continue
            norm_target = remove_accents(target_col)
            if col_norm == norm_target:
                matched_targets.add(target_col)
                break
            for alias in aliases:
                norm_alias = remove_accents(alias)
                if col_norm == norm_alias or re.search(r'\b' + re.escape(norm_alias) + r'\b', col_norm):
                    matched_targets.add(target_col)
                    break
    return len(matched_targets)



def get_all_candidate_files(directory: Path = None) -> List[Path]:
    """Lấy danh sách tất cả file dữ liệu nguồn hợp lệ trong thư mục (.xlsx, .xls, .xlsm, .csv, .tsv, .txt, .xml, .json)."""
    if directory is None:
        directory = Path.cwd()
    valid_extensions = {".xlsx", ".xls", ".xlsm", ".csv", ".tsv", ".txt", ".xml", ".json", ".jsonl"}
    ignored_names = {
        "requirements.txt", "readme.md", "dockerfile", "docker-compose.yml",
        ".gitignore", "license", "setup.py"
    }
    candidate_files = []
    for f in directory.iterdir():
        if f.is_file() and f.suffix.lower() in valid_extensions:
            # Bỏ qua file tạm Excel (~$) và các file kết quả xuất ra, cũng như file cấu hình dự án
            if (
                not f.name.startswith("~$")
                and not f.name.startswith(".")
                and not f.name.startswith("Processed_")
                and not f.name.startswith("ClickAI_")
                and f.name.lower() not in ignored_names
            ):
                candidate_files.append(f)
    candidate_files.sort(key=lambda x: x.stat().st_mtime, reverse=True)
    return candidate_files


def find_input_file(specified_path: str = None) -> Path:
    """Tự động tìm kiếm file dữ liệu nguồn (.xlsx, .xls, .xlsm, .csv, .tsv, .txt, .xml, .json)."""
    if specified_path:
        p = Path(specified_path)
        if p.exists() and p.is_file():
            return p
        raise FileNotFoundError(f"Không tìm thấy file nguồn chỉ định: {specified_path}")

    candidate_files = get_all_candidate_files()
    if not candidate_files:
        raise FileNotFoundError(
            "Không tìm thấy file đầu vào hợp lệ (.xlsx, .xls, .csv, .tsv, .txt, .xml, .json) nào trong thư mục hiện tại!\n"
            "Vui lòng đặt file sản phẩm vào cùng thư mục hoặc dùng tham số --input <tên_file>."
        )

    selected = candidate_files[0]
    if len(candidate_files) > 1:
        print(f"[*] Phát hiện {len(candidate_files)} file đầu vào. Chọn file mới nhất: {selected.name}")

        print(f"    (Danh sách: {[f.name for f in candidate_files]})")
    return selected


def load_smart_excel(file_path: Path) -> pd.DataFrame:
    """
    Đọc file Excel thông minh:
    - Hỗ trợ nhiều Sheet: Tự động phát hiện và gộp các Sheet liên quan chứa sản phẩm.
    - Tự động dò tìm dòng Tiêu đề (Header Row): Tránh trường hợp các dòng đầu là logo/báo cáo/dòng trống.
    """
    engine = "openpyxl" if file_path.suffix.lower() in [".xlsx", ".xlsm"] else None
    try:
        xls = pd.ExcelFile(file_path, engine=engine)
    except Exception:
        xls = pd.ExcelFile(file_path)

    sheet_names = xls.sheet_names
    print(f"[*] File có {len(sheet_names)} sheet: {sheet_names}")

    product_sheets = []  # Danh sách sheet chứa sản phẩm: (sheet, header_row, score)

    for sheet in sheet_names:
        try:
            sample_df = pd.read_excel(xls, sheet_name=sheet, header=None, nrows=15)
        except Exception:
            continue
        if sample_df.empty:
            continue

        best_r = 0
        max_score = 0
        for r_idx in range(min(15, len(sample_df))):
            row_vals = sample_df.iloc[r_idx].dropna().tolist()
            score = count_matched_product_columns(row_vals)
            if score > max_score:
                max_score = score
                best_r = r_idx

        sheet_norm = remove_accents(sheet)
        bonus = 2 if any(k in sheet_norm for k in ["san pham", "hang hoa", "product", "item", "data", "danh sach"]) else 0
        total_score = max_score + bonus

        # Nếu sheet có chứa ít nhất 2 cột trùng với từ điển sản phẩm
        if max_score >= 2 or (max_score >= 1 and bonus > 0):
            product_sheets.append((sheet, best_r, total_score))

    # Nếu không tìm thấy sheet nào chuyên biệt, đọc sheet đầu tiên
    if not product_sheets:
        print(f"[*] Không phát hiện sheet chuyên biệt, đọc sheet đầu tiên: '{sheet_names[0]}'")
        df = pd.read_excel(xls, sheet_name=sheet_names[0])
        return df.dropna(how="all")

    # Nếu chỉ có 1 sheet chứa sản phẩm
    if len(product_sheets) == 1:
        best_sheet, best_header_row, _ = product_sheets[0]
        print(f"[*] Đã chọn Sheet: '{best_sheet}' (Dòng tiêu đề: hàng {best_header_row + 1})")
        df = pd.read_excel(xls, sheet_name=best_sheet, header=best_header_row)
        return df.dropna(how="all")

    # Nếu có NHIỀU sheet liên quan chứa sản phẩm (ví dụ: từng danh mục ở mỗi sheet)
    print(f"[*] Phát hiện {len(product_sheets)} sheet liên quan chứa sản phẩm: {[s[0] for s in product_sheets]}")
    dfs = []
    for sheet, h_row, _ in product_sheets:
        sub_df = pd.read_excel(xls, sheet_name=sheet, header=h_row).dropna(how="all")
        if not sub_df.empty:
            dfs.append(sub_df)

    if dfs:
        combined_df = pd.concat(dfs, ignore_index=True)
        print(f"[✓] Đã tự động gộp thành công {len(dfs)} sheet liên quan (Tổng cộng: {len(combined_df)} dòng dữ liệu).")
        return combined_df

    return pd.read_excel(xls, sheet_name=sheet_names[0]).dropna(how="all")



def load_smart_csv(file_path: Path) -> pd.DataFrame:
    """
    Đọc file CSV thông minh:
    - Thử nhiều encoding (utf-8-sig, utf-8, windows-1252, cp1258, latin1)
    - Tự động nhận diện dấu phân cách (, hoặc ;)
    - Tự động bỏ qua các dòng tiêu đề báo cáo/dòng trống phía trên
    """
    encodings = ["utf-8-sig", "utf-8", "windows-1252", "cp1258", "latin1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc, errors="ignore") as f:
                lines = [f.readline() for _ in range(25)]
            lines = [l for l in lines if l]
            if not lines:
                continue

            for sep in [",", ";", "\t"]:
                best_r = 0
                max_score = 0
                for r_idx, line in enumerate(lines):
                    if not line.strip():
                        continue
                    parts = [p.strip().strip('"').strip("'") for p in line.split(sep)]
                    score = count_matched_product_columns(parts)
                    if score > max_score:
                        max_score = score
                        best_r = r_idx

                skip_rows = best_r if max_score >= 1 else 0
                try:
                    df = pd.read_csv(file_path, encoding=enc, sep=sep, skiprows=skip_rows)
                    if not df.empty and len(df.columns) > 1:
                        print(f"[*] Đọc CSV thành công với encoding='{enc}', sep='{sep}', dòng tiêu đề: hàng {skip_rows + 1}")
                        return df.dropna(how="all")
                except Exception:
                    continue
        except Exception:
            continue
    raise ValueError(f"Không thể đọc file văn bản CSV/TSV '{file_path.name}'. Vui lòng kiểm tra định dạng file.")


def load_smart_xml(file_path: Path) -> pd.DataFrame:
    """
    Đọc và chuẩn hóa dữ liệu từ file XML:
    1. Hỗ trợ Excel 2003 XML (SpreadsheetML - phổ biến từ MISA, Fast, KiotViet, ERP).
    2. Hỗ trợ Google Merchant Feed, RSS Product Catalog, WooCommerce XML, Generic XML.
    """
    print(f"[*] Đang phân tích cấu trúc file XML: {file_path.name}")
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
    except Exception as e:
        raise ValueError(f"Không thể đọc cấu trúc file XML: {e}")

    # Nhận diện SpreadsheetML (Excel XML 2003)
    root_tag = root.tag.lower()
    is_spreadsheetml = "spreadsheet" in root_tag or any("spreadsheet" in str(v).lower() for v in root.attrib.values())
    if not is_spreadsheetml:
        for elem in root.iter():
            tag_clean = elem.tag.split("}")[-1].lower()
            if tag_clean in ["worksheet", "table"]:
                is_spreadsheetml = True
                break

    if is_spreadsheetml:
        print(f"[*] Nhận diện định dạng Excel XML (SpreadsheetML / Office XML).")
        worksheets = [elem for elem in root.iter() if elem.tag.split("}")[-1].lower() == "worksheet"]
        if not worksheets:
            worksheets = [root]

        all_sheet_dfs = []
        for ws in worksheets:
            ws_name = ws.attrib.get("{urn:schemas-microsoft-com:office:spreadsheet}Name") or ws.attrib.get("Name") or "Sheet"
            tables = [elem for elem in ws.iter() if elem.tag.split("}")[-1].lower() == "table"]
            target_table = tables[0] if tables else ws
            rows_elems = [elem for elem in target_table if elem.tag.split("}")[-1].lower() == "row"]

            raw_rows = []
            for r_elem in rows_elems:
                cells = [c for c in r_elem if c.tag.split("}")[-1].lower() == "cell"]
                row_dict = {}
                col_cursor = 0
                for cell in cells:
                    # Kiểm tra ss:Index
                    index_val = None
                    for k, v in cell.attrib.items():
                        if k.split("}")[-1].lower() == "index":
                            try:
                                index_val = int(v) - 1
                            except ValueError:
                                pass
                            break
                    if index_val is not None:
                        col_cursor = index_val

                    # Lấy text trong <Data>
                    data_elem = next((c for c in cell if c.tag.split("}")[-1].lower() == "data"), None)
                    text = data_elem.text.strip() if (data_elem is not None and data_elem.text) else ""
                    row_dict[col_cursor] = text
                    col_cursor += 1

                if row_dict:
                    max_c = max(row_dict.keys())
                    row_list = [row_dict.get(c, "") for c in range(max_c + 1)]
                    if any(str(x).strip() for x in row_list):
                        raw_rows.append(row_list)

            if not raw_rows:
                continue

            # Dò tìm dòng Header Row tối ưu
            best_r = 0
            max_score = 0
            for r_idx in range(min(20, len(raw_rows))):
                score = count_matched_product_columns(raw_rows[r_idx])
                if score > max_score:
                    max_score = score
                    best_r = r_idx

            if max_score >= 1 and len(raw_rows) > best_r + 1:
                header_cols = raw_rows[best_r]
                seen_cols = {}
                clean_headers = []
                for idx, col in enumerate(header_cols):
                    c_name = str(col).strip() if col else f"Cot_{idx+1}"
                    if not c_name:
                        c_name = f"Cot_{idx+1}"
                    if c_name in seen_cols:
                        seen_cols[c_name] += 1
                        c_name = f"{c_name}_{seen_cols[c_name]}"
                    else:
                        seen_cols[c_name] = 1
                    clean_headers.append(c_name)

                data_rows = raw_rows[best_r + 1:]
                col_len = len(clean_headers)
                normalized_data = []
                for r in data_rows:
                    if len(r) < col_len:
                        r = r + [""] * (col_len - len(r))
                    else:
                        r = r[:col_len]
                    if any(str(x).strip() for x in r):
                        normalized_data.append(r)

                sheet_df = pd.DataFrame(normalized_data, columns=clean_headers)
                print(f"[*] Sheet '{ws_name}': Đọc thành công {len(sheet_df)} dòng từ hàng header {best_r + 1}.")
                all_sheet_dfs.append(sheet_df)
            else:
                if len(raw_rows) > 1:
                    sheet_df = pd.DataFrame(raw_rows[1:], columns=raw_rows[0])
                    all_sheet_dfs.append(sheet_df)

        if all_sheet_dfs:
            combined_df = pd.concat(all_sheet_dfs, ignore_index=True) if len(all_sheet_dfs) > 1 else all_sheet_dfs[0]
            print(f"[i] Đã nạp thành công từ SpreadsheetML: {len(combined_df)} dòng sản phẩm, {len(combined_df.columns)} cột.")
            return combined_df

    # 2. Xử lý XML dạng Feeds / RSS / WooCommerce / Generic XML
    target_tags = ["item", "product", "entry", "record", "row", "san_pham", "hang_hoa", "item_data"]
    items = []

    for tag in target_tags:
        found = root.findall(f".//{tag}")
        if found:
            items = found
            print(f"[*] Tìm thấy {len(items)} phần tử sản phẩm trong thẻ '<{tag}>'")
            break

    if not items:
        tag_counts = {}
        for elem in root.iter():
            tag_name = elem.tag.split("}")[-1]
            tag_counts[tag_name] = tag_counts.get(tag_name, 0) + 1

        candidate_tags = [t for t, c in tag_counts.items() if c >= 2]
        for t in candidate_tags:
            elems = [e for e in root.iter() if e.tag.split("}")[-1] == t and len(e) > 0]
            if len(elems) >= 2:
                items = elems
                print(f"[*] Tự động nhận diện cấu trúc lặp sản phẩm trong thẻ '<{t}>' ({len(items)} sản phẩm)")
                break

    if not items:
        raise ValueError(f"Không tìm thấy danh sách sản phẩm lặp lại trong file XML '{file_path.name}'.")

    rows = []
    for item in items:
        row_dict = {}
        for attr_k, attr_v in item.attrib.items():
            row_dict[attr_k] = attr_v

        for child in item:
            tag_name = child.tag
            if "}" in tag_name:
                ns, local_tag = tag_name.split("}", 1)
                if "google" in ns.lower() or "base" in ns.lower():
                    clean_tag = f"g:{local_tag}"
                else:
                    clean_tag = local_tag
            else:
                clean_tag = tag_name

            text_val = child.text.strip() if child.text else ""
            if text_val:
                row_dict[clean_tag] = text_val
            elif len(child) > 0:
                sub_texts = [c.text.strip() for c in child if c.text and c.text.strip()]
                if sub_texts:
                    row_dict[clean_tag] = ", ".join(sub_texts)

        if row_dict:
            rows.append(row_dict)

    if not rows:
        raise ValueError(f"File XML '{file_path.name}' không có nội dung dữ liệu sản phẩm nào.")

    df = pd.DataFrame(rows)
    print(f"[i] Đã nạp thành công từ XML: {len(df)} sản phẩm, {len(df.columns)} trường dữ liệu.")
    return df


def load_smart_json(file_path: Path) -> pd.DataFrame:
    """
    Đọc dữ liệu từ file JSON / JSONL (Shopify Export, REST API JSON, Database dump,...).
    Tự động chuẩn hóa mảng JSON hoặc bóc tách từ các khóa lồng nhau (products, items, data).
    """
    print(f"[*] Đang đọc file JSON: {file_path.name}")
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        first_chunk = f.read(200).strip()
        f.seek(0)

        # Kiểm tra JSON Lines (nhiều object JSON trên nhiều dòng)
        if file_path.suffix.lower() == ".jsonl" or ("\n{" in first_chunk and not first_chunk.startswith("[")):
            rows = []
            for line in f:
                line = line.strip()
                if line:
                    try:
                        rows.append(json.loads(line))
                    except Exception:
                        continue
            if rows:
                print(f"[*] Nhận diện định dạng JSON Lines: {len(rows)} bản ghi.")
                return pd.json_normalize(rows)

        f.seek(0)
        try:
            content = json.load(f)
        except Exception as e:
            raise ValueError(f"File JSON không hợp lệ: {e}")

    if isinstance(content, list):
        return pd.json_normalize(content)

    elif isinstance(content, dict):
        product_keys = ["products", "items", "data", "records", "result", "rows", "san_pham"]
        for k in product_keys:
            if k in content and isinstance(content[k], list) and len(content[k]) > 0:
                print(f"[*] Tìm thấy danh sách sản phẩm trong khóa '{k}' ({len(content[k])} sản phẩm).")
                return pd.json_normalize(content[k])

        for k, v in content.items():
            if isinstance(v, list) and len(v) > 0 and isinstance(v[0], dict):
                print(f"[*] Tự động nhận diện danh sách sản phẩm trong khóa '{k}' ({len(v)} sản phẩm).")
                return pd.json_normalize(v)

        return pd.json_normalize([content])

    raise ValueError(f"Không nhận diện được cấu trúc dữ liệu sản phẩm trong file JSON '{file_path.name}'.")


def load_dataframe(file_path: Path) -> pd.DataFrame:
    """Đọc dữ liệu từ mọi định dạng: Excel (.xlsx, .xls, .xlsm), CSV/TSV (.csv, .tsv, .txt), XML (.xml), JSON (.json, .jsonl)."""
    ext = file_path.suffix.lower()
    print(f"[*] Đang đọc file: {file_path.name} (Định dạng: {ext})")

    if ext in [".xlsx", ".xls", ".xlsm"]:
        return load_smart_excel(file_path)
    elif ext in [".csv", ".tsv", ".txt"]:
        return load_smart_csv(file_path)
    elif ext == ".xml":
        return load_smart_xml(file_path)
    elif ext in [".json", ".jsonl"]:
        return load_smart_json(file_path)
    else:
        raise ValueError(
            f"Định dạng file không được hỗ trợ: '{ext}'.\n"
            f"Hệ thống hỗ trợ: Excel (.xlsx, .xls, .xlsm), CSV/TSV (.csv, .tsv, .txt), XML (.xml), JSON (.json, .jsonl)."
        )




def clean_numeric_value(val):
    """Làm sạch giá trị số (loại bỏ ký tự đ, VND, dấu phẩy ngăn cách hàng nghìn)."""
    if pd.isna(val) or val is None or str(val).strip() == "":
        return None
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    # Loại bỏ đơn vị tiền tệ hoặc chữ cái
    s = re.sub(r'[^\d.,\-]', '', s)
    if not s:
        return None
    # Xử lý dấu phân cách hàng nghìn dạng 100,000 hoặc 100.000
    if ',' in s and '.' in s:
        # 1,234.56 -> 1234.56
        s = s.replace(',', '')
    elif ',' in s and '.' not in s:
        parts = s.split(',')
        if len(parts[-1]) == 3:  # 100,000 -> 100000
            s = s.replace(',', '')
        else:  # 12,5 (số thập phân ở VN/EU)
            s = s.replace(',', '.')
    elif '.' in s and ',' not in s:
        parts = s.split('.')
        # Nếu có nhiều dấu chấm (100.000.000) hoặc 100.000 thì là ngăn cách hàng nghìn
        if len(parts) > 2 or (len(parts) == 2 and len(parts[-1]) == 3):
            s = s.replace('.', '')
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return val


def inspect_column_content(series: pd.Series) -> Optional[str]:
    """
    Phân tích nội dung thực tế của các ô dữ liệu trong cột (Content Heuristics)
    để tự động nhận diện ý nghĩa cột ngay cả khi tên tiêu đề không liên quan.
    """
    clean_series = series.dropna().astype(str).str.strip()
    clean_series = clean_series[~clean_series.isin(["", "nan", "None", "null"])]
    if len(clean_series) == 0:
        return None

    sample = clean_series.head(40).tolist()
    n = len(sample)

    # 1. Nhận diện Link hình ảnh (.jpg, .png, .webp, /images/, cdn...)
    img_matches = sum(
        1 for s in sample 
        if any(s.lower().endswith(ext) for ext in [".jpg", ".jpeg", ".png", ".webp", ".gif", ".svg"])
        or ("http" in s.lower() and any(k in s.lower() for k in ["/image", "/img", "/photo", "cdn.", "/thumb"]))
    )
    if img_matches / n >= 0.5:
        return "Link hình ảnh"

    # 2. Nhận diện Link sản phẩm (URL website nhưng không phải link ảnh)
    url_matches = sum(
        1 for s in sample 
        if s.lower().startswith("http://") or s.lower().startswith("https://") or s.lower().startswith("www.")
    )
    if url_matches / n >= 0.5:
        return "Link sản phẩm"

    # 3. Phân biệt dữ liệu dạng Văn bản (Text) và dạng Số / Mã (Numeric/Code)
    text_items = [s for s in sample if len(re.findall(r'[a-zA-ZÀ-ỹ]', s)) >= 3 and not s.startswith("http")]
    is_text_column = (len(text_items) / n >= 0.6)

    if is_text_column:
        # Nhận diện Xuất xứ (Quốc gia sản xuất)
        origin_keywords = ["vietnam", "viet nam", "trung quoc", "china", "han quoc", "korea", "nhat ban", "japan", "my", "usa", "thai lan", "duc", "phap"]
        origin_matches = sum(1 for s in sample if remove_accents(s) in origin_keywords or "made in" in s.lower())
        if origin_matches / n >= 0.4:
            return "Xuất xứ"

        # Nhận diện Tùy chọn giá / Biến thể (Size, màu sắc)
        option_keywords = ["size", "den", "trang", "do", "xanh", "vang", "xl", "xxl", "xs", "sm", "md", "lg"]
        opt_matches = sum(1 for s in sample if any(w in remove_accents(s).split() for w in option_keywords))
        if opt_matches / n >= 0.4:
            return "Tùy chọn giá"

        # Nhận diện Mô tả (Chứa HTML hoặc văn bản câu dài > 120 ký tự)
        has_html = any("<p>" in s.lower() or "<br" in s.lower() or "<div" in s.lower() for s in sample)
        avg_len = sum(len(s) for s in sample) / n
        if has_html or avg_len > 120:
            return "Mô tả"

        # Nhận diện Tên sản phẩm (Văn bản có từ 2-25 từ, độ dài 8-110 ký tự)
        avg_words = sum(len(s.split()) for s in text_items) / len(text_items)
        if 2 <= avg_words <= 25 and 8 <= avg_len <= 110:
            return "Tên sản phẩm"

    else:
        # Cột dạng Mã hoặc Số
        # Nhận diện Mã SKU / Mã hàng
        sku_pattern = re.compile(r'^[A-Za-z0-9_\-\.\/]{3,35}$')
        sku_matches = sum(1 for s in sample if sku_pattern.match(s) and not s.isdigit())
        if sku_matches / n >= 0.6:
            return "Mã SKU"

        # Nhận diện Giá tiền hoặc Tồn kho
        has_currency_sign = any(any(c in s.lower() for c in ["đ", "vnd", "vnđ", "usd", "$"]) for s in sample)
        num_values = []
        for s in sample:
            v = clean_numeric_value(s)
            if isinstance(v, (int, float)):
                num_values.append(v)

        if len(num_values) / n >= 0.6:
            avg_val = sum(num_values) / len(num_values)
            if has_currency_sign or (avg_val >= 1000 and all(v >= 0 for v in num_values)):
                return "Giá bán"
            elif all(float(v).is_integer() and 0 <= v <= 100000 for v in num_values):
                return "Tồn kho"

    return None




def build_column_mapping(source_df: pd.DataFrame) -> Tuple[Dict[str, str], List[str]]:
    """
    Tạo bảng ánh xạ thông minh từ DataFrame nguồn sang 16 cột đích của ClickAI:
    - Bước 1: Khớp chính xác tên cột
    - Bước 2: Khớp qua từ điển bí danh (Synonyms & Aliases)
    - Bước 3: Phân tích nội dung thực tế trong từng ô (Content Heuristics)
    - Bước 4: Thu thập các cột tính năng/thuộc tính chưa map để gộp vào Mô tả
    Trả về: (mapping_dict, unmapped_feature_columns)
    """
    source_columns = [str(c) for c in source_df.columns if str(c) != "_source_sheet_name"]
    mapping = {}
    used_source_cols = set()

    # Chuẩn hóa tên cột nguồn để so sánh
    normalized_sources = []
    for col in source_columns:
        norm = remove_accents(col)
        normalized_sources.append((col, norm))

    # Bước 1: Khớp chính xác hoàn toàn (Exact Match)
    for target_col in CLICKAI_COLUMNS:
        norm_target = remove_accents(target_col)
        for orig_col, norm_col in normalized_sources:
            if orig_col in used_source_cols:
                continue
            if norm_col == norm_target:
                mapping[target_col] = orig_col
                used_source_cols.add(orig_col)
                break

    # Bước 2: Khớp qua từ điển bí danh (Synonyms & Aliases)
    for target_col, aliases in SYNONYMS_MAP.items():
        if target_col in mapping:
            continue
        sorted_aliases = sorted(aliases, key=len, reverse=True)
        matched_col = None
        for alias in sorted_aliases:
            norm_alias = remove_accents(alias)
            for orig_col, norm_col in normalized_sources:
                if orig_col in used_source_cols:
                    continue
                # Khớp hoàn toàn
                if norm_col == norm_alias:
                    matched_col = orig_col
                    break
                # Khớp cụm từ (không áp dụng cho từ đơn chung chung để tránh nhận nhầm)
                generic_single_words = {"product", "name", "item", "type", "code", "gia", "price", "anh", "image"}
                if norm_alias not in generic_single_words and (" " in norm_alias or len(norm_alias) >= 6):
                    if re.search(r'\b' + re.escape(norm_alias) + r'\b', norm_col):
                        matched_col = orig_col
                        break
            if matched_col:
                break

        if matched_col:
            mapping[target_col] = matched_col
            used_source_cols.add(matched_col)

    # Bước 3: Phân tích nội dung thực tế của các cột chưa map (Content Inspection)
    for orig_col in source_columns:
        if orig_col in used_source_cols:
            continue
        predicted_target = inspect_column_content(source_df[orig_col])
        if predicted_target and predicted_target not in mapping:
            print(f" [🔍 Nội dung ô tính] Phát hiện cột '{orig_col}' chứa dữ liệu '{predicted_target}' -> Ánh xạ tự động!")
            mapping[predicted_target] = orig_col
            used_source_cols.add(orig_col)

    # Bước 4: Thu thập các cột thuộc tính / tính năng còn lại chưa map (để ghép vào Mô tả)
    unmapped_feature_cols = []
    ignored_keywords = ["unnamed", "stt", "id", "index", "_source_sheet_name"]
    for orig_col in source_columns:
        if orig_col not in used_source_cols:
            norm_c = remove_accents(orig_col)
            if not any(k in norm_c for k in ignored_keywords):
                clean_s = source_df[orig_col].dropna().astype(str).str.strip()
                if len(clean_s) > 0:
                    unmapped_feature_cols.append(orig_col)

    # Kiểm tra tính hợp lệ tối thiểu của file dữ liệu
    core_cols = {"Tên sản phẩm", "Mã SKU", "Giá bán"}
    matched_core = [c for c in core_cols if c in mapping]
    if len(mapping) == 0 or not matched_core:
        sample_cols = [str(c) for c in source_columns[:8]]
        raise ValueError(
            "Không nhận diện được cột sản phẩm hợp lệ trong file (cần tối thiểu cột 'Tên sản phẩm', 'Mã SKU' hoặc 'Giá bán'). "
            f"Các cột tìm thấy: {sample_cols}. Vui lòng kiểm tra lại cấu trúc file!"
        )

    return mapping, unmapped_feature_cols


def process_dataframe(source_df: pd.DataFrame, mapping: dict, unmapped_feature_cols: list = None) -> pd.DataFrame:
    """Tạo DataFrame chuẩn ClickAI 16 cột, gộp tính năng chưa map vào Mô tả và điền mặc định."""
    target_data = {}

    print("\n" + "=" * 60)
    print("BẢNG KẾT QUẢ ÁNH XẠ CỘT (COLUMN MAPPING):")
    print("=" * 60)
    for target_col in CLICKAI_COLUMNS:
        if target_col in mapping:
            src_col = mapping[target_col]
            target_data[target_col] = source_df[src_col].copy()
            print(f" [✓] {target_col:<20} <--- [Cột gốc: '{src_col}']")
        else:
            target_data[target_col] = pd.Series([None] * len(source_df))
            print(f" [!] {target_col:<20} <--- (Không có trong file gốc, tạo cột trống)")

    processed_df = pd.DataFrame(target_data)

    # Tự động gộp các thuộc tính / tính năng chưa map vào cột 'Mô tả'
    if unmapped_feature_cols:
        print(f"\n[*] Phát hiện {len(unmapped_feature_cols)} cột tính năng/thông số phụ: {unmapped_feature_cols}")
        print("    -> Tự động gộp nội dung các tính năng này vào cột 'Mô tả' để AI đọc được đầy đủ...")

        for idx, row in source_df.iterrows():
            feature_lines = []
            for col in unmapped_feature_cols:
                val = row[col]
                if pd.notna(val) and str(val).strip() not in ("", "nan", "None"):
                    feature_lines.append(f"- {col}: {str(val).strip()}")
            if feature_lines:
                specs_text = "📌 Thông số & Tính năng bổ sung:\n" + "\n".join(feature_lines)
                existing_desc = processed_df.at[idx, "Mô tả"]
                if pd.isna(existing_desc) or str(existing_desc).strip() in ("", "nan", "None"):
                    processed_df.at[idx, "Mô tả"] = specs_text
                else:
                    processed_df.at[idx, "Mô tả"] = f"{str(existing_desc).strip()}\n\n{specs_text}"

    # Chuẩn hóa kiểu dữ liệu số cho các cột tính toán
    numeric_cols = ["Giá bán", "Giá vốn", "Tồn kho", "Thuế (%)", "Chi phí trực tiếp"]
    for col in numeric_cols:
        processed_df[col] = processed_df[col].apply(clean_numeric_value)

    # Điền giá trị mặc định nếu trường bị trống / NaN
    print("\n[*] Áp dụng giá trị mặc định cho dữ liệu trống:")
    for col, default_val in DEFAULT_VALUES.items():
        count_missing = processed_df[col].isna().sum() + (processed_df[col] == "").sum()
        if count_missing > 0:
            processed_df[col] = processed_df[col].fillna(default_val)
            processed_df[col] = processed_df[col].replace("", default_val)
            print(f"    - '{col}': Đã điền '{default_val}' cho {count_missing} dòng trống.")
        else:
            print(f"    - '{col}': Đầy đủ dữ liệu (không cần điền mặc định).")

    # Giới hạn ký tự cell tối đa của Excel để tránh lỗi file
    for col in CLICKAI_COLUMNS:
        processed_df[col] = processed_df[col].apply(
            lambda x: str(x)[:EXCEL_CELL_LIMIT] if isinstance(x, str) and len(str(x)) > EXCEL_CELL_LIMIT else x
        )

    return processed_df



def is_multisheet_export(file_path: Path) -> bool:
    """Nhận diện file export website nhiều sheet (có 'Sản phẩm' VÀ ('Biến thể' hoặc 'Thông số'))."""
    if file_path.suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        names = set(wb.sheetnames)
        wb.close()
        return "Sản phẩm" in names and ("Biến thể" in names or "Thông số" in names)
    except Exception:
        return False



def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "None")


def _first(*vals):
    """Trả về giá trị không rỗng đầu tiên."""
    for v in vals:
        if not _blank(v):
            return v
    return None


def _to_int(v, default=0):
    v = clean_numeric_value(v)
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


EXCEL_CELL_LIMIT = 32000
ORIGIN_SPEC_NAMES = ["Xuất xứ", "Nơi sản xuất", "Sản xuất", "Xuất Xứ"]


def convert_multisheet_export(file_path: Path) -> pd.DataFrame:
    """
    Chuyển file export nhiều sheet (Sản phẩm / Biến thể / Thông số) sang 16 cột ClickAI.
    - Sản phẩm simple            -> 1 dòng.
    - Sản phẩm variable có biến thể -> mỗi biến thể 1 dòng (Tùy chọn giá = tên phân loại).
    """
    print(f"[*] Phát hiện file export nhiều sheet: {file_path.name}")
    xls = pd.ExcelFile(file_path, engine="openpyxl")
    sp = pd.read_excel(xls, "Sản phẩm")
    bt = pd.read_excel(xls, "Biến thể") if "Biến thể" in xls.sheet_names else pd.DataFrame()
    ts = pd.read_excel(xls, "Thông số") if "Thông số" in xls.sheet_names else pd.DataFrame()
    print(f"[i] Sheet 'Sản phẩm': {len(sp)} SP | 'Biến thể': {len(bt)} dòng | 'Thông số': {len(ts)} dòng")

    # Xuất xứ theo SKU (lấy từ sheet Thông số)
    origin_map = {}
    if not ts.empty and {"SKU", "Thông số", "Giá trị"} <= set(ts.columns):
        o = ts[ts["Thông số"].astype(str).str.strip().isin(ORIGIN_SPEC_NAMES)].dropna(subset=["Giá trị"])
        origin_map = o.drop_duplicates("SKU").set_index("SKU")["Giá trị"].astype(str).str.strip().to_dict()

    # Nhóm biến thể theo SKU cha (chỉ lấy biến thể đang hoạt động)
    variants = {}
    if not bt.empty and "SKU" in bt.columns:
        if "is_active" in bt.columns:
            bt = bt[bt["is_active"].astype(str).str.strip() != "Không"]
        for sku, grp in bt.groupby("SKU", sort=False):
            variants[sku] = grp

    def g(row, col):
        return row[col] if col in row.index else None

    rows = []
    n_variant_rows = 0
    for _, p in sp.iterrows():
        sku = g(p, "SKU")
        desc = _first(g(p, "Mô tả ngắn"), g(p, "Mô tả chi tiết (text)"), g(p, "Mô tả"))
        if desc is not None:
            desc = str(desc).strip()[:EXCEL_CELL_LIMIT]
        base = {
            "Tên sản phẩm": g(p, "Tên sản phẩm"),
            "Tùy chọn giá": None,
            "Mã SKU": sku,
            "Danh mục": _first(g(p, "Danh mục chính"), g(p, "Đường dẫn danh mục chính")),
            "Giá bán": clean_numeric_value(g(p, "Giá bán (VND)")),
            "Giá vốn": None,  # File nguồn không có giá vốn ("Giá gốc" là giá niêm yết trước giảm)
            "Tồn kho": _to_int(g(p, "Tồn kho (SL)")),
            "Đơn vị tính": _first(g(p, "Đơn vị tính"), g(p, "ĐVT cơ bản")),
            "Tiền tệ": "VND",
            "Thuế (%)": 0,
            "Chi phí trực tiếp": 0,
            "Chu kỳ thanh toán": "Một lần",
            "Xuất xứ": origin_map.get(sku),
            "Mô tả": desc,
            "Link hình ảnh": g(p, "Ảnh chính"),
            "Link sản phẩm": g(p, "Link trên website"),
        }

        vgrp = variants.get(sku)
        if vgrp is not None and len(vgrp) > 0:
            for _, v in vgrp.iterrows():
                r = dict(base)
                r["Tùy chọn giá"] = _first(g(v, "options_label"), g(v, "Thuộc tính biến thể"))
                r["Mã SKU"] = _first(g(v, "variant_sku"), g(v, "version_sku"), sku)
                price = clean_numeric_value(_first(g(v, "Giá áp dụng"), g(v, "price")))
                r["Giá bán"] = price if price is not None else base["Giá bán"]
                r["Tồn kho"] = _to_int(g(v, "stock_quantity"))
                r["Link hình ảnh"] = _first(g(v, "image"), base["Link hình ảnh"])
                rows.append(r)
                n_variant_rows += 1
        else:
            rows.append(base)

    df = pd.DataFrame(rows, columns=CLICKAI_COLUMNS)

    # Chuẩn hóa ĐVT viết hoa chữ đầu
    df["Đơn vị tính"] = df["Đơn vị tính"].apply(lambda x: None if _blank(x) else str(x).strip().capitalize())

    # Áp mặc định lần cuối cho chắc chắn
    for col, val in DEFAULT_VALUES.items():
        df[col] = df[col].apply(lambda x: val if _blank(x) else x)

    print(f"[i] Kết quả: {len(df)} dòng ({n_variant_rows} dòng từ biến thể, "
          f"{len(df) - n_variant_rows} dòng sản phẩm đơn).")
    print(f"[i] Có Xuất xứ: {df['Xuất xứ'].notna().sum()} dòng | "
          f"Có ĐVT: {df['Đơn vị tính'].notna().sum()} dòng | Giá vốn: để trống (nguồn không có).")
    return df


def save_to_formatted_excel(df: pd.DataFrame, output_path: Path):
    """Xuất DataFrame ra file Excel với định dạng chuyên nghiệp, dễ đọc."""
    print(f"\n[*] Đang lưu kết quả ra file: {output_path.name}")
    
    # Ghi file tạm bằng pandas
    df.to_excel(output_path, index=False, engine="openpyxl")

    # Mở lại bằng openpyxl để định dạng thẩm mỹ
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    ws.title = "Sản Phẩm ClickAI"

    # Định dạng Header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Xanh navy đậm
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(CLICKAI_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Định dạng các dòng dữ liệu
    data_font = Font(name="Calibri", size=10)
    regular_align_left = Alignment(horizontal="left", vertical="center")
    regular_align_right = Alignment(horizontal="right", vertical="center")
    regular_align_center = Alignment(horizontal="center", vertical="center")

    right_align_cols = {"Giá bán", "Giá vốn", "Tồn kho", "Thuế (%)", "Chi phí trực tiếp"}
    center_align_cols = {"Đơn vị tính", "Tiền tệ", "Chu kỳ thanh toán", "Xuất xứ", "Mã SKU"}

    total_rows = ws.max_row
    # Tối ưu hóa hiệu năng: Định dạng border & font cho tối đa 300 dòng đầu để đảm bảo tốc độ phản hồi < 1 giây trên Cloud
    format_limit_rows = min(total_rows, 300)

    for row_idx in range(2, format_limit_rows + 1):
        ws.row_dimensions[row_idx].height = 20
        # Tạo hiệu ứng sọc ngựa vằn (zebra striping) nhẹ
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if is_even else None

        for col_idx in range(1, len(CLICKAI_COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            col_name = CLICKAI_COLUMNS[col_idx - 1]
            if col_name in right_align_cols:
                cell.alignment = regular_align_right
                if isinstance(cell.value, (int, float)):
                    if col_name == "Thuế (%)":
                        cell.number_format = '#,##0.0' if isinstance(cell.value, float) else '#,##0'
                    else:
                        cell.number_format = '#,##0'
            elif col_name in center_align_cols:
                cell.alignment = regular_align_center
            else:
                cell.alignment = regular_align_left

    # Tự động căn chỉnh độ rộng cột thông minh
    for col_idx in range(1, len(CLICKAI_COLUMNS) + 1):
        col_letter = get_column_letter(col_idx)
        header_text = str(ws.cell(row=1, column=col_idx).value or '')
        max_len = len(header_text)

        # Kiểm tra độ dài 50 dòng đầu tiên để tối ưu hiệu năng
        for row in range(2, min(ws.max_row + 1, 52)):
            val = ws.cell(row=row, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))

        # Giới hạn chiều rộng hợp lý (tối thiểu 12, tối đa 40)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 40), 12)

    # Bật hiển thị gridlines
    ws.views.sheetView[0].showGridLines = True

    wb.save(output_path)
    print(f"[✓] Đã xuất thành công: {output_path.resolve()}")


def process_product_file(input_path: Union[str, Path], output_path: Optional[Union[str, Path]] = None) -> Path:
    """
    Hàm chuẩn hóa chính dùng cho cả chạy lệnh dòng lệnh (CLI) và tích hợp API Server.
    Nhận vào đường dẫn file (.xlsx, .xls, .csv), xuất ra file chuẩn ClickAI 16 cột.
    """
    input_file = Path(input_path)
    if not input_file.exists():
        raise FileNotFoundError(f"Không tìm thấy file nguồn: {input_path}")

    if output_path is None:
        output_file = input_file.parent / f"Processed_{input_file.stem}.xlsx"
    else:
        output_file = Path(output_path)

    if is_multisheet_export(input_file):
        processed_df = convert_multisheet_export(input_file)
    else:
        source_df = load_dataframe(input_file)
        row_count, col_count = source_df.shape
        print(f"[i] Đã nạp thành công: {row_count} dòng sản phẩm, {col_count} cột dữ liệu.")

        mapping, unmapped_feature_cols = build_column_mapping(source_df)
        processed_df = process_dataframe(source_df, mapping, unmapped_feature_cols)

    save_to_formatted_excel(processed_df, output_file)
    return output_file


def main():
    parser = argparse.ArgumentParser(
        description="ClickAI Product Mapper - Tự động chuẩn hóa dữ liệu sản phẩm sang chuẩn ClickAI."
    )
    parser.add_argument(
        "-i", "--input",
        type=str,
        default=None,
        help="Đường dẫn file đầu vào (.xlsx, .xls, .csv). Nếu bỏ trống sẽ tự động quét tất cả file trong thư mục."
    )
    parser.add_argument(
        "-o", "--output",
        type=str,
        default=None,
        help="Đường dẫn file Excel đầu ra."
    )
    args = parser.parse_args()

    print("=" * 65)
    print("      CLICKAI PRODUCT MAPPER & SMART BATCH NORMALIZER")
    print("=" * 65)

    current_dir = Path.cwd()

    try:
        # Trường hợp 1: Người dùng chỉ định cụ thể 1 file
        if args.input:
            input_file = Path(args.input)
            output_file = Path(args.output) if args.output else current_dir / f"Processed_{input_file.stem}.xlsx"
            print(f"[▶ ĐANG XỬ LÝ] File chỉ định: {input_file.name}...")
            result_path = process_product_file(input_file, output_file)
            shutil.copyfile(result_path, current_dir / "ClickAI_Product_Processed.xlsx")
            print("=" * 65)
            print("HOÀN THÀNH XỬ LÝ DỮ LIỆU THÀNH CÔNG!")
            print(f"- File xuất: {result_path.resolve()}")
            print("=" * 65)
            return

        # Trường hợp 2: Chạy tự động toàn bộ thư mục (User gõ: python3 product_mapper_local.py)
        candidate_files = get_all_candidate_files(current_dir)
        if not candidate_files:
            print("[!] Không tìm thấy file dữ liệu nào (.xlsx, .xls, .csv) trong thư mục để xử lý.")
            print("    Vui lòng copy file sản phẩm vào thư mục này rồi chạy lại lệnh.")
            return

        print(f"[*] Tìm thấy {len(candidate_files)} file dữ liệu trong thư mục.")
        print("-" * 65)

        count_processed = 0
        count_skipped = 0
        latest_result_path = None

        for input_file in candidate_files:
            output_file = current_dir / f"Processed_{input_file.stem}.xlsx"

            # Tự động nhận diện file đã làm rồi: file kết quả đã có và mới hơn file nguồn
            if output_file.exists() and output_file.stat().st_mtime >= input_file.stat().st_mtime:
                print(f"[⏩ BỎ QUA] '{input_file.name}' (Đã xử lý trước đó -> '{output_file.name}')")
                count_skipped += 1
                continue

            print(f"\n[▶ ĐANG XỬ LÝ] '{input_file.name}'...")
            try:
                result_path = process_product_file(input_file, output_file)
                latest_result_path = result_path
                print(f"[✓ HOÀN THÀNH] '{input_file.name}' -> '{output_file.name}'")
                count_processed += 1
            except Exception as e:
                print(f"[✗ LỖI] Không thể xử lý file '{input_file.name}': {e}")

        # Đồng bộ file ClickAI_Product_Processed.xlsx với file mới nhất
        if latest_result_path and latest_result_path.exists():
            shutil.copyfile(latest_result_path, current_dir / "ClickAI_Product_Processed.xlsx")

        print("\n" + "=" * 65)
        print(f"TỔNG KẾT BATCH:")
        print(f"  - File mới đã xử lý: {count_processed}")
        print(f"  - File cũ đã bỏ qua: {count_skipped}")
        if latest_result_path:
            print(f"  - File chuẩn ClickAI mới nhất: ClickAI_Product_Processed.xlsx")
        print("=" * 65)

    except Exception as e:
        print(f"\n[LỖI] Có lỗi xảy ra trong quá trình xử lý: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()


